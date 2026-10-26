from django.core.paginator import Paginator
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.db.models import Count
from .models import Student, Attendance
from .forms import StudentForm

from io import BytesIO
from openpyxl.styles import Font, PatternFill, Alignment
from django.utils.http import urlsafe_base64_encode, urlsafe_base64_decode

from django.contrib.auth.tokens import default_token_generator
from django.utils.encoding import force_bytes
from django.urls import reverse
from django.core.mail import send_mail
from django.conf import settings

from django.contrib.auth import login, logout
from django.contrib.auth.forms import UserCreationForm, AuthenticationForm
from django.contrib.auth.decorators import login_required, user_passes_test
import csv
from django.http import HttpResponse
from openpyxl import Workbook

from openpyxl import load_workbook
from django.core.validators import validate_email
from django.core.exceptions import ValidationError
from django.db import transaction

from functools import wraps
from django.http import HttpResponseForbidden


from functools import wraps
from django.http import HttpResponseForbidden

from django.contrib.auth.models import User
from django.contrib.auth import get_user_model

from django.contrib.auth.forms import PasswordChangeForm
from django.contrib.auth import update_session_auth_hash

from datetime import date

from django.db.models import Count, Q


def admin_required(view_func):
    @wraps(view_func)
    def wrapper(request, *args, **kwargs):

        if not request.user.is_authenticated:
            return redirect('login')

        if not request.user.is_staff:
            return render(
                request,
                'students/403.html',
                status=403
            )

        return view_func(request, *args, **kwargs)

    return wrapper


@login_required
def student_list(request):
    students = Student.objects.all()

    search_query = request.GET.get('search', '')
    department = request.GET.get('department', '')

    if search_query:
        students = students.filter(name__icontains=search_query)

    if department:
        students = students.filter(department__iexact=department)

    departments = Student.objects.values_list(
        'department', flat=True
    ).distinct().order_by('department')

    paginator = Paginator(students, 5)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    return render(request, 'students/student_list.html', {
        'students': page_obj,
        'page_obj': page_obj,
        'search_query': search_query,
        'selected_department': department,
        'departments': departments,
    })
@admin_required
@login_required
def dashboard(request):
    total_students = Student.objects.count()

    total_departments = Student.objects.values(
        'department'
    ).distinct().count()

    recent_students = Student.objects.order_by(
        '-created_at'
    )[:5]

    department_counts = Student.objects.values(
        'department'
    ).annotate(
        total=Count('id')
    ).order_by('department')

    context = {
        'total_students': total_students,
        'total_departments': total_departments,
        'recent_students': recent_students,
        'department_counts': list(department_counts),
    }

    return render(request, 'students/dashboard.html', context)

@admin_required
@login_required
def add_student(request):
    if request.method == 'POST':
        form = StudentForm(request.POST)

        if form.is_valid():
            form.save()
            messages.success(request, 'Student added successfully!')
            return redirect('student_list')
    else:
        form = StudentForm()

    return render(request, 'students/add_student.html', {
        'form': form
    })

@admin_required
@login_required
def edit_student(request, student_id):
    student = get_object_or_404(Student, id=student_id)

    if request.method == 'POST':
        form = StudentForm(request.POST, instance=student)

        if form.is_valid():
            form.save()
            messages.success(request, 'Student updated successfully!')
            return redirect('student_list')
    else:
        form = StudentForm(instance=student)

    return render(request, 'students/edit_student.html', {
        'form': form,
        'student': student
    })

@admin_required
@login_required
def delete_student(request, student_id):
    student = get_object_or_404(Student, id=student_id)

    if request.method == 'POST':
        student.delete()
        messages.success(request, 'Student deleted successfully!')

    return redirect('student_list')


def register(request):
    if request.method == 'POST':
        form = UserCreationForm(request.POST)

        if form.is_valid():
            user = form.save(commit=False)

            # Get the email entered by the user
            email = request.POST.get('email', '').strip().lower()

            if not email:
                form.add_error(None, 'Email address is required.')
                return render(
                    request,
                    'students/register.html',
                    {'form': form}
                )

            # Validate the email address
            try:
                validate_email(email)
            except ValidationError:
                form.add_error(None, 'Enter a valid email address.')
                return render(
                    request,
                    'students/register.html',
                    {'form': form}
                )

            # Prevent duplicate email addresses
            User = get_user_model()

            if User.objects.filter(email__iexact=email).exists():
                form.add_error(None, 'This email address is already registered.')
                return render(
                    request,
                    'students/register.html',
                    {'form': form}
                )

            # Create an inactive regular user
            user.email = email
            user.is_active = False
            user.is_staff = False
            user.is_superuser = False
            user.save()

            # Generate the verification link
            uid = urlsafe_base64_encode(force_bytes(user.pk))
            token = default_token_generator.make_token(user)

            verification_path = reverse(
                'verify_email',
                kwargs={'uidb64': uid, 'token': token}
            )

            verification_url = request.build_absolute_uri(
                verification_path
            )

            # Send the verification email
            try:
                send_mail(
                    subject='Verify Your Student Management System Account',
                    message=(
                        f'Hello {user.username},\n\n'
                        'Please click the link below to verify your email '
                        'address and activate your account:\n\n'
                        f'{verification_url}\n\n'
                        'If you did not create this account, you can '
                        'ignore this email.'
                    ),
                    from_email=settings.DEFAULT_FROM_EMAIL,
                    recipient_list=[user.email],
                    fail_silently=False,
                )
            except Exception:
                user.delete()
                messages.error(
                    request,
                    'Unable to send the verification email. '
                    'Please try registering again later.'
                )
                return redirect('register')

            messages.success(
                request,
                'Registration successful! Check your email, including '
                'the Spam folder, to verify your account.'
            )
            return redirect('login')

    else:
        form = UserCreationForm()

    return render(request, 'students/register.html', {'form': form})

def verify_email(request, uidb64, token):
    User = get_user_model()

    try:
        uid = urlsafe_base64_decode(uidb64).decode()
        user = User.objects.get(pk=uid)
    except (TypeError, ValueError, OverflowError, User.DoesNotExist):
        user = None

    if user is not None and default_token_generator.check_token(user, token):
        if user.is_active:
            messages.info(request, 'This account is already active.')
        else:
            user.is_active = True
            user.save(update_fields=['is_active'])
            messages.success(
                request,
                'Email verified successfully! You can now log in.'
            )

        return redirect('login')

    messages.error(
        request,
        'This verification link is invalid or has expired.'
    )
    return redirect('register')


def user_login(request):
    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '')

        user = User.objects.filter(username=username).first()

        if user and not user.is_active:
            messages.error(
                request,
                "Your account is inactive or your email is not verified. "
                "Please verify your email or contact the administrator."
            )
            return redirect('login')

        form = AuthenticationForm(
            request,
            data={
                'username': username,
                'password': password
            }
        )

        if form.is_valid():
            user = form.get_user()
            login(request, user)
            return redirect('student_list')

        messages.error(request, "Invalid username or password.")

    else:
        form = AuthenticationForm(request)

    return render(request, 'students/login.html', {'form': form})


def user_logout(request):
    if request.method == 'POST':
        logout(request)
        return redirect('login')

    return redirect('student_list')

# Export student data to CSV
@admin_required
@login_required
def export_students_csv(request):
    students = Student.objects.all().order_by('id')

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="students.csv"'

    writer = csv.writer(response)

    writer.writerow([
        'ID', 'Name', 'Email', 'Age', 'Department', 'Date Added'
    ])

    for student in students:
        writer.writerow([
            student.id,
            student.name,
            student.email,
            student.age,
            student.department,
            student.created_at.strftime('%Y-%m-%d %H:%M:%S')
            if student.created_at else ''
        ])

    return response


# Export student data to Excel
@admin_required
@login_required
def export_students_excel(request):
    students = Student.objects.all().order_by('id')

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Students'

    worksheet.append([
        'ID', 'Name', 'Email', 'Age', 'Department', 'Date Added'
    ])

    for student in students:
        worksheet.append([
            student.id,
            student.name,
            student.email,
            student.age,
            student.department,
            student.created_at.strftime('%Y-%m-%d %H:%M:%S')
            if student.created_at else ''
        ])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="students.xlsx"'

    workbook.save(response)

    return response
@admin_required
@login_required
def import_students(request):

    if request.method == 'POST':

        excel_file = request.FILES.get('excel_file')

        # Check whether a file was uploaded
        if not excel_file:
            messages.error(request, "Please select an Excel file.")
            return redirect('import_students')

        # Check file extension
        if not excel_file.name.lower().endswith('.xlsx'):
            messages.error(request, "Please upload an .xlsx file.")
            return redirect('import_students')

        # Limit the upload to 5 MB
        if excel_file.size > 5 * 1024 * 1024:
            messages.error(request, "File size must not exceed 5 MB.")
            return redirect('import_students')

        try:
            workbook = load_workbook(
                excel_file,
                read_only=True,
                data_only=True
            )
            sheet = workbook.active

            rows = sheet.iter_rows(values_only=True)

            # Read the first row as column headers
            headers = next(rows, None)

            if not headers:
                workbook.close()
                messages.error(request, "The Excel file is empty.")
                return redirect('import_students')

            # Normalize header names
            header_map = {
                str(value).strip().lower(): index
                for index, value in enumerate(headers)
                if value is not None
            }

            required_headers = [
                'name', 'email', 'age', 'department'
            ]

            missing_headers = [
                header for header in required_headers
                if header not in header_map
            ]

            if missing_headers:
                workbook.close()
                messages.error(
                    request,
                    "Missing columns: " + ", ".join(missing_headers)
                )
                return redirect('import_students')

            students_to_create = []
            imported_emails = set()
            errors = []

            # Read each student row
            for row_number, row in enumerate(rows, start=2):

                # Skip completely empty rows
                if not any(
                    value is not None and str(value).strip()
                    for value in row
                ):
                    continue

                try:
                    name = row[header_map['name']]
                    email = row[header_map['email']]
                    age = row[header_map['age']]
                    department = row[header_map['department']]

                    # Check required fields
                    if any(
                        value is None or not str(value).strip()
                        for value in [name, email, age, department]
                    ):
                        raise ValueError(
                            "All four fields are required."
                        )

                    name = str(name).strip()
                    email = str(email).strip().lower()
                    department = str(department).strip()

                    # Validate email
                    validate_email(email)

                    # Validate age
                    if isinstance(age, bool):
                        raise ValueError("Age must be a positive integer.")

                    if isinstance(age, str):
                        if not age.strip().isdigit():
                            raise ValueError(
                                "Age must be a positive integer."
                            )
                        age = int(age)
                    elif isinstance(age, (int, float)):
                        if isinstance(age, float) and not age.is_integer():
                            raise ValueError("Age must be a whole number.")
                        age = int(age)
                    else:
                        raise ValueError(
                            "Age must be a positive integer."
                        )

                    if age <= 0:
                        raise ValueError(
                            "Age must be greater than zero."
                        )

                    # Check duplicate emails within the file
                    if email in imported_emails:
                        raise ValueError(
                            "This email appears more than once in the file."
                        )

                    imported_emails.add(email)

                    students_to_create.append(
                        Student(
                            name=name,
                            email=email,
                            age=age,
                            department=department
                        )
                    )

                except (ValueError, ValidationError, IndexError) as error:
                    errors.append(
                        f"Row {row_number}: {error}"
                    )

            workbook.close()

            if not students_to_create and not errors:
                messages.error(
                    request,
                    "No student records were found in the Excel file."
                )
                return redirect('import_students')

            # Check whether emails already exist in the database
            existing_emails = {
                email.lower()
                for email in Student.objects.values_list(
                    'email', flat=True
                )
            }

            for student in students_to_create:
                if student.email in existing_emails:
                    errors.append(
                        f"Email '{student.email}' already exists "
                        "in the database."
                    )

            # If any error exists, import nothing
            if errors:
                for error in errors:
                    messages.error(request, error)

                messages.error(
                    request,
                    "Import cancelled. No students were added."
                )
                return redirect('import_students')

            # Save all records together
            with transaction.atomic():
                Student.objects.bulk_create(students_to_create)

            messages.success(
                request,
                f"Successfully imported {len(students_to_create)} students!"
            )

            return redirect('student_list')

        except Exception as error:
            messages.error(
                request,
                "Unable to read the Excel file. Please check that "
                "it is a valid .xlsx workbook."
            )
            return redirect('import_students')

    return render(request, 'students/import_students.html')


@login_required
def student_detail(request, student_id):
    student = get_object_or_404(Student, id=student_id)

    return render(
        request,
        'students/student_detail.html',
        {'student': student}
    )

@admin_required
@login_required
def user_management(request):
    users = User.objects.all().order_by('id')

    return render(
        request,
        'students/user_management.html',
        {'users': users}
    )

@admin_required
@login_required
def change_user_role(request, user_id):

    if request.method == 'POST':
        account = get_object_or_404(User, id=user_id)

        # Prevent changing your own role
        if account.id == request.user.id:
            messages.error(
                request,
                "You cannot change your own role."
            )
            return redirect('user_management')

        # Prevent changing the main superuser account
        if account.is_superuser:
            messages.error(
                request,
                "You cannot change the role of a superuser."
            )
            return redirect('user_management')

        if account.is_staff:
            account.is_staff = False
            messages.success(
                request,
                f"{account.username} is now a Regular User."
            )
        else:
            account.is_staff = True
            messages.success(
                request,
                f"{account.username} is now an Admin."
            )

        account.save()

    return redirect('user_management')

@admin_required
@login_required
def toggle_user_status(request, user_id):

    if request.method == 'POST':
        account = get_object_or_404(User, id=user_id)

        # Prevent changing your own account status
        if account.id == request.user.id:
            messages.error(
                request,
                "You cannot deactivate your own account."
            )
            return redirect('user_management')

        # Protect the main superuser account
        if account.is_superuser:
            messages.error(
                request,
                "You cannot change the status of a superuser."
            )
            return redirect('user_management')

        # Activate or deactivate the account
        account.is_active = not account.is_active
        account.save()

        if account.is_active:
            messages.success(
                request,
                f"{account.username} has been activated."
            )
        else:
            messages.success(
                request,
                f"{account.username} has been deactivated."
            )

    return redirect('user_management')

@login_required
def change_password(request):

    if request.method == 'POST':
        form = PasswordChangeForm(
            user=request.user,
            data=request.POST
        )

        if form.is_valid():
            user = form.save()

            # Keep the user logged in after changing password
            update_session_auth_hash(request, user)

            messages.success(
                request,
                "Your password has been changed successfully."
            )

            return redirect('student_list')

    else:
        form = PasswordChangeForm(user=request.user)

    return render(
        request,
        'students/change_password.html',
        {'form': form}
    )


from datetime import date
from django.db import transaction


@login_required
def take_attendance(request):
    departments = Student.objects.values_list(
        'department', flat=True
    ).distinct().order_by('department')

    department = request.GET.get(
        'department',
        request.POST.get('department', '')
    ).strip()

    class_name = request.GET.get(
        'class_name',
        request.POST.get('class_name', '')
    ).strip()

    attendance_date = request.GET.get(
        'date',
        request.POST.get('date', date.today().isoformat())
    )

    # Validate the selected date
    try:
        selected_date = date.fromisoformat(attendance_date)
    except (ValueError, TypeError):
        selected_date = date.today()
        attendance_date = selected_date.isoformat()

    # Load classes belonging to the selected department
    classes = (
        Student.objects.filter(department=department)
        .values_list('class_name', flat=True)
        .distinct()
        .order_by('class_name')
        if department
        else []
    )

    # Load students belonging to the selected department and class
    students = Student.objects.none()

    if department and class_name:
        students = Student.objects.filter(
            department=department,
            class_name=class_name
        ).order_by('name')

    # Save attendance when the form is submitted
    if request.method == 'POST':
        if not department or not class_name:
            messages.error(
                request,
                'Please select a department and class.'
            )
        elif not students.exists():
            messages.error(
                request,
                'No students found in the selected class.'
            )
        else:
            with transaction.atomic():
                for student in students:
                    status = request.POST.get(
                        f'status_{student.id}', 'Absent'
                    )

                    if status not in [
                        'Present', 'Absent', 'On Leave'
                    ]:
                        status = 'Absent'

                    Attendance.objects.update_or_create(
                        student=student,
                        date=selected_date,
                        defaults={
                            'status': status,
                            'marked_by': request.user
                        }
                    )

            messages.success(
                request,
                'Attendance saved successfully!'
            )

        # Redirect while preserving the selected filters
        from urllib.parse import urlencode

        query_string = urlencode({
            'department': department,
            'class_name': class_name,
            'date': attendance_date,
        })

        return redirect(
            f"{reverse('take_attendance')}?{query_string}"
        )

    # Retrieve previously saved attendance for this date
    existing_records = {
        record.student_id: record.status
        for record in Attendance.objects.filter(
            student__in=students,
            date=selected_date
        )
    }

    # Make each student's saved status available in the template
    for student in students:
        student.attendance_status = existing_records.get(
            student.id, 'Absent'
        )

    return render(
        request,
        'students/take_attendance.html',
        {
            'departments': departments,
            'classes': classes,
            'students': students,
            'selected_department': department,
            'selected_class': class_name,
            'attendance_date': attendance_date,
        }
    )
    
    
@login_required
def attendance_history(request):
    students = Student.objects.all().order_by('name')

    student_id = request.GET.get('student', '').strip()
    start_date = request.GET.get('start_date', '').strip()
    end_date = request.GET.get('end_date', '').strip()

    records = Attendance.objects.select_related('student').all()

    selected_student = None
    valid_student_id = None

    if student_id:
        try:
            valid_student_id = int(student_id)
            selected_student = students.filter(
                id=valid_student_id
            ).first()

            if selected_student:
                records = records.filter(
                    student=selected_student
                )
            else:
                valid_student_id = None
                student_id = ''
        except (ValueError, TypeError):
            student_id = ''

    if start_date:
        try:
            date.fromisoformat(start_date)
            records = records.filter(date__gte=start_date)
        except ValueError:
            start_date = ''

    if end_date:
        try:
            date.fromisoformat(end_date)
            records = records.filter(date__lte=end_date)
        except ValueError:
            end_date = ''

    if start_date and end_date and start_date > end_date:
        records = Attendance.objects.none()

    records = records.order_by('-date', 'student__name')

    summary = records.aggregate(
        total=Count('id'),
        present=Count('id', filter=Q(status='Present')),
        absent=Count('id', filter=Q(status='Absent')),
        on_leave=Count('id', filter=Q(status='On Leave')),
    )

    total = summary['total']
    present = summary['present']

    percentage = round((present / total) * 100, 2) if total else 0

    return render(
        request,
        'students/attendance_history.html',
        {
            'students': students,
            'records': records,
            'selected_student': selected_student,
            'student_id': student_id,
            'start_date': start_date,
            'end_date': end_date,
            'total': total,
            'present': present,
            'absent': summary['absent'],
            'on_leave': summary['on_leave'],
            'percentage': percentage,
        }
    )


@user_passes_test(
    lambda user: user.is_authenticated and user.is_staff,
    login_url='login'
)
def export_attendance_excel(request):
    # Get optional date filters from the URL
    start_date = request.GET.get('start_date', '').strip()
    end_date = request.GET.get('end_date', '').strip()

    records = Attendance.objects.select_related(
        'student', 'marked_by'
    ).all().order_by('-date', 'student__name')

    # Validate and apply the starting date
    if start_date:
        try:
            date.fromisoformat(start_date)
            records = records.filter(date__gte=start_date)
        except ValueError:
            return HttpResponse(
                'Invalid start date.',
                status=400
            )

    # Validate and apply the ending date
    if end_date:
        try:
            date.fromisoformat(end_date)
            records = records.filter(date__lte=end_date)
        except ValueError:
            return HttpResponse(
                'Invalid end date.',
                status=400
            )

    if start_date and end_date and start_date > end_date:
        return HttpResponse(
            'Start date cannot be later than end date.',
            status=400
        )

    # Create Excel workbook
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Attendance Report'

    headers = [
        'Attendance ID',
        'Student Name',
        'Email',
        'Department',
        'Class',
        'Date',
        'Status',
        'Marked By',
    ]

    worksheet.append(headers)

    # Style the header row
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(
            fill_type='solid',
            fgColor='2563EB'
        )
        cell.alignment = Alignment(horizontal='center')

    # Add attendance records
    for record in records:
        worksheet.append([
            record.id,
            record.student.name,
            record.student.email,
            record.student.department,
            record.student.class_name,
            record.date.strftime('%d-%m-%Y'),
            record.status,
            record.marked_by.username
            if record.marked_by else 'N/A',
        ])

    # Adjust column widths
    for column in worksheet.columns:
        max_length = max(
            len(str(cell.value or ''))
            for cell in column
        )
        worksheet.column_dimensions[
            column[0].column_letter
        ].width = min(max_length + 3, 35)

    worksheet.freeze_panes = 'A2'
    worksheet.auto_filter.ref = worksheet.dimensions

    # Return the Excel file as a download
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            'application/vnd.openxmlformats-officedocument.'
            'spreadsheetml.sheet'
        )
    )
    response['Content-Disposition'] = (
        'attachment; filename="attendance_report.xlsx"'
    )

    return response




